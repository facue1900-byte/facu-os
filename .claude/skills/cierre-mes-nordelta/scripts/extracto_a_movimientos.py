#!/usr/bin/env python3
"""
Extracto Macro -> Movimientos del Master Plan.

Lee el PDF del extracto (CC Especial en Pesos 4-452-0960512147-9), identifica cada
crédito y débito, y arma las filas en el orden de columnas de la pestaña
Movimientos: Fecha, Tipo, Medio, Local, Categoria, Monto, Moneda, Observaciones.

    # sólo mirar
    extracto_a_movimientos.py "Julio 2026.pdf" --facturas "…/Julio 2026"
    # escribir de verdad
    extracto_a_movimientos.py "Julio 2026.pdf" --facturas "…/Julio 2026" --escribir

`--facturas` es lo que evita el trabajo a mano. Varios débitos del extracto
salen como «N/D Transf. MacrOnline E-set D/T» sin CUIT ni nombre: la glosa no
dice a quién se le pagó. En vez de que alguien abra las facturas del mes y las
cruce de a una, el script indexa los PDF de la carpeta, saca CUIT y total de
cada uno y **matchea por importe** (el banco redondea al peso, por eso la
tolerancia de $1). El CUIT decide la categoría.

`--escribir` es el único que toca la planilla, y **se planta si queda un solo
renglón sin resolver**: un movimiento sin categoría entra igual a Movimientos,
no rompe nada y desaparece de todos los SUMIFS. Es exactamente la falla que
deja una expensa más barata sin que nadie se entere.

El dedupe compara (año, mes, monto) contra lo que ya está cargado con Medio
«Banco»: Facu a veces carga un movimiento con fecha distinta a la del extracto.

La columna I (Mes) es ArrayFormula: NO se emite y no hay que tocarla.
"""

import sys
import re
import datetime
import unicodedata

import argparse
import os
import glob

import fitz

sys.path.insert(0, "/Users/Facu/facu-os/execution")

CUENTA = "4-452-0960512147-9"
MASTER_PLAN = "1ATiNBHCukPYPn9-poP1HO4SlfsDu5pGXsLz-JvW-IQs"

# CUIT del proveedor -> (nombre corto, categoría de Movimientos).
# El CUIT sale del nombre del archivo de la factura (convención de AFIP,
# `CUIT_pv_tipo_nro.pdf`) o del texto del PDF. Un CUIT que no esté acá deja el
# movimiento SIN resolver a propósito: se agrega el renglón, no se adivina.
CUIT_CATEGORIA = {
    "27373389973": ("Rhino (Sánchez Yanina)", "Sueldo Mantenimiento Gastronomia"),
    "30690741691": ("Andersen materiales",    "Inversiones"),
    "30517431431": ("Transportes Olivos",     "Retiro de Residuos"),
    "30718796136": ("Estudio Giaccio",        "Contador"),
    "20124767173": ("Oliva (corralón)",       "Inversiones"),
    # Agregados con el extracto de agosto 2026, cada uno contra su comprobante
    # de pago de la subcarpeta «Pagos/».
    "30707845682": ("Redes y Servicios",      "Redes y Servicios"),
    # Giaccio Mariano Jesús, persona física: factura el balance anual aparte del
    # abono mensual que factura Giaccio Business Consulting SRL (30718796136).
    # Los dos son el contador.
    "20081049980": ("Giaccio Mariano Jesús",  "Contador"),
    # GEVGE S.R.L., el proveedor que factura como «Electro 927»: materiales
    # eléctricos e iluminación para la obra. Identificado el 09/09/2026 leyendo
    # con OCR las facturas escaneadas del grupo de WhatsApp «Facturas Paseo
    # Nordelta» — sus PDF no tienen texto, así que el CUIT no salía de ningún lado.
    "30710063474": ("GEVGE (Electro 927)",    "Inversiones"),
    # Agregados con el extracto de septiembre 2026, cada uno con su precedente en
    # Movimientos: Annie ya iba a Marketing (ene-26) y los matafuegos a Mejoras (ene-26).
    "27447142959": ("Hoffer Annie (diseño redes)", "Marketing / Redes Sociales"),
    "20205123149": ("Matafuegos Yany",        "Mejoras"),
}

# Glosa del extracto -> (Local, Categoria). Local para ingresos, Categoria para egresos.
# El orden importa: gana la primera que matchea.
REGLAS = [
    (r"SUSHINOR",                        "Fabric",                None),
    (r"RODOLFO SRL|30716281457",         "Bigg",                  None),
    # Identificados en el extracto de julio 2026 contra las facturas. El CUIT
    # viaja en la glosa de las TRANSF, así que se matchean solos.
    (r"20286590994",                     None, "Sueldo Administracion (MB)"),   # Matías Barbagrigia
    # Oliva Gustavo Alberto — corralón de Tigre, materiales entregados en obra.
    # Se propone Inversiones porque hasta hoy siempre fue obra; si una compra es
    # gasto operativo hay que cambiarla a mano ANTES de cargarla.
    (r"20124767173",                     None, "Inversiones"),
    # Transportes Olivos («TODSE»): ya estaba en CUIT_CATEGORIA, pero eso sólo
    # sirve cuando el débito se matchea contra una factura de la carpeta. En
    # agosto 2026 el CUIT vino en la glosa («TRANSF 30517431431 VAR»), que no
    # pasa por ahí. Con la regla se resuelve por los dos caminos.
    (r"30517431431",                     None, "Retiro de Residuos"),
    # GEVGE («Electro 927»): igual que Oliva y Andersen, materiales de obra.
    (r"30710063474",                     None, "Inversiones"),
    # Transferencia entre cuentas propias para cubrir el resumen de la VISA en
    # la CC Bancaria. NO es plata que se queda adentro: del otro lado sale a la
    # tarjeta. En junio 2026 se cargó como Inversiones («cubre VISA luces») y se
    # sigue ese precedente. Si el consumo del mes no fue obra, hay que cambiarla.
    (r"TRANSF AUT SDO MISMO TIT",        None, "Inversiones"),
    (r"RET\.? ING\.? BRUTOS SIRCREB",    None, "Ingresos Brutos"),
    (r"DBCR 25413",                      None, "Gastos bancarios"),
    (r"Comision Trf|COMISION",           None, "Gastos bancarios"),
    (r"IMP\.? AFIP|VEP",                 None, "Impuestos"),
    (r"MANTENIMIENTO|SERV\. CTA",        None, "Gastos bancarios"),
]

MESES = {1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
         7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre",
         12: "diciembre"}

LINEA = re.compile(r"^\s*(\d{2}/\d{2}/\d{2})\s+(.*?)\s+((?:\d{1,3}\.)*\d{1,3},\d{2})"
                   r"\s+((?:-?\d{1,3}\.)*-?\d{1,3},\d{2})\s*$")


def plata(s):
    return float(s.replace(".", "").replace(",", "."))


def sin_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def ar(n):
    """1234567.8 -> '1.234.567,80', el formato del resumen mensual."""
    return f"{n:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def texto_movimientos_export(doc):
    """Rearma el export «movimientos» del home banking con el layout del resumen.

    Cuando el resumen mensual no llegó, el Macro deja bajar los movimientos de la
    cuenta (`mov_caja_de_ahorros_*.pdf`, emitido a pedido). Trae lo mismo pero en
    otro formato: números en inglés (1,234.56), débito y crédito en COLUMNAS sin
    saldo por renglón, y sólo un «SALDO FINAL DEL DIA». La columna se lee de la
    coordenada x del importe. El saldo de cada renglón se reconstruye y se exige
    que caiga exacto en cada saldo diario, en el final y en los totales de la
    cabecera: si una sola cosa no da, no se devuelve nada.
    """
    num = re.compile(r"^\d{1,3}(?:,\d{3})*\.\d{2}$")
    en = lambda s: float(s.replace(",", ""))  # noqa: E731
    cab = None
    filas = []  # (fecha, desc, debito, credito) o ("DIA", saldo)
    for pagina in doc:
        palabras = pagina.get_text("words")
        heads = {w[4]: w for w in palabras if w[4] in ("DEBITOS", "CREDITOS", "SALDO")
                 and 280 < w[1] < 305}
        if cab is None:
            # Cabecera: SALDO INICIAL · CREDITOS · DEBITOS · IVA · SALDO FINAL, en ese orden de x.
            tot = sorted((w for w in palabras if num.match(w[4]) and w[1] < 260),
                         key=lambda w: w[0])
            if len(tot) == 5:
                cab = [en(w[4]) for w in tot]
        # Débito y crédito salen alineados a la derecha: el borde de la columna es
        # el punto medio entre los encabezados.
        corte_dc = (heads["DEBITOS"][2] + heads["CREDITOS"][0]) / 2 + 30
        corte_cs = (heads["CREDITOS"][2] + heads["SALDO"][0]) / 2 + 30
        renglones = {}
        for w in palabras:
            if w[1] > 305:
                renglones.setdefault(round(w[1]), []).append(w)
        for y in sorted(renglones):
            ws = sorted(renglones[y], key=lambda w: w[0])
            if not re.match(r"^\d{2}/\d{2}/\d{4}$", ws[0][4]):
                continue
            fecha = ws[0][4]
            texto_l = " ".join(w[4] for w in ws[1:] if not num.match(w[4]))
            montos = [w for w in ws[1:] if num.match(w[4])]
            if len(montos) != 1:
                sys.exit(f"Renglón {fecha} '{texto_l}' con {len(montos)} importes.")
            m = montos[0]
            if "SALDO FINAL" in texto_l:
                filas.append(("DIA", fecha, en(m[4])))
            elif m[2] <= corte_dc:
                filas.append((fecha, texto_l, en(m[4]), 0.0))
            elif m[2] <= corte_cs:
                filas.append((fecha, texto_l, 0.0, en(m[4])))
            else:
                sys.exit(f"Renglón {fecha} '{texto_l}': importe en la columna de saldo.")
        final_txt = re.search(r"SALDO FINAL AL D[IÍ]A (\d{2}/\d{2}/\d{4}) :\s*([\d,.]+)",
                              pagina.get_text())
        if final_txt:
            fecha_final, saldo_final = final_txt.group(1), en(final_txt.group(2))
    if cab is None or not filas:
        sys.exit("El export de movimientos no trae la cabecera de saldos o está vacío.")
    inicial, cred_tot, deb_tot, _iva, final = cab
    saldo, lineas = inicial, []
    deb = cred = 0.0
    for f in filas:
        if f[0] == "DIA":
            if abs(saldo - f[2]) > 0.005:
                sys.exit(f"El saldo reconstruido al {f[1]} da {ar(saldo)} y el banco "
                         f"dice {ar(f[2])}: alguna columna se leyó mal.")
            continue
        fecha, desc, d, c = f
        saldo = round(saldo - d + c, 2)
        deb, cred = deb + d, cred + c
        dd, mm, aaaa = fecha.split("/")
        lineas.append(f"  {dd}/{mm}/{aaaa[2:]} {desc}   {ar(d or c)}   {ar(saldo)}")
    for nombre, calc, banco in (("débitos", deb, deb_tot), ("créditos", cred, cred_tot),
                                ("saldo final", saldo, final), ("saldo al pie", saldo, saldo_final)):
        if abs(calc - banco) > 0.005:
            sys.exit(f"Export de movimientos: {nombre} da {ar(calc)} y la cabecera dice {ar(banco)}.")
    dd, mm, aaaa = fecha_final.split("/")
    previo = datetime.date(int(aaaa), int(mm), 1) - datetime.timedelta(days=1)
    return "\n".join(
        [f"CUENTA CORRIENTE ESPECIAL EN PESOS NRO.: {CUENTA}",
         f"   SALDO ULTIMO EXTRACTO AL {previo:%d/%m/%Y}   {ar(inicial)}"]
        + lineas
        + [f"   SALDO FINAL AL DIA {fecha_final}   {ar(saldo)}"])


def texto_pdf(pdf_path):
    """Texto del extracto en el layout del resumen mensual, venga como venga."""
    doc = fitz.open(pdf_path)
    texto = "\n".join(p.get_text() for p in doc)
    if f"CUENTA CORRIENTE ESPECIAL EN PESOS NRO.: {CUENTA}" in texto:
        return texto
    if CUENTA in texto and "SALDO FINAL DEL DIA" in texto:
        return texto_movimientos_export(doc)
    return texto


def bloque_cuenta(texto):
    """Aísla el detalle de la cuenta en pesos: el extracto trae tres cuentas."""
    ini = texto.find(f"CUENTA CORRIENTE ESPECIAL EN PESOS NRO.: {CUENTA}")
    if ini < 0:
        sys.exit(f"No encontré la cuenta {CUENTA} en el PDF.")
    resto = texto[ini + 10:]
    fin = resto.find("CUENTA CORRIENTE BANCARIA NRO.")
    return resto[:fin] if fin > 0 else resto


def clasificar(desc):
    d = sin_tildes(desc).upper()
    for patron, local, categoria in REGLAS:
        if re.search(patron, d, re.I):
            return local, categoria
    return None, None


def parsear(pdf_path):
    bloque = bloque_cuenta(texto_pdf(pdf_path))

    saldo_prev, movs = None, []
    for linea in bloque.split("\n"):
        if "SALDO ULTIMO EXTRACTO" in linea:
            m = re.search(r"((?:\d{1,3}\.)*\d{1,3},\d{2})\s*$", linea)
            if m:
                saldo_prev = plata(m.group(1))
            continue
        m = LINEA.match(linea)
        if not m or saldo_prev is None:
            continue
        fecha_s, desc, monto_s, saldo_s = m.groups()
        monto, saldo = plata(monto_s), plata(saldo_s)
        # El extracto no separa débito de crédito por columna de forma confiable:
        # lo deduzco de cómo se movió el saldo.
        tipo = "Ingreso" if saldo > saldo_prev else "Egreso"
        saldo_prev = saldo
        d, mth, y = (int(x) for x in fecha_s.split("/"))
        local, categoria = clasificar(desc)
        movs.append({
            "fecha": datetime.date(2000 + y, mth, d),
            "tipo": tipo,
            "local": local if tipo == "Ingreso" else None,
            "categoria": categoria if tipo == "Egreso" else None,
            "monto": monto,
            "obs": " ".join(desc.split()),
        })
    return movs


def ya_cargados(xlsx):
    """(año, mes, monto) de los movimientos de Banco que ya están en el sheet.

    El match es por mes y no por día: Facu a veces carga un movimiento con
    fecha distinta a la del extracto (ej. el sueldo del 26/6 quedó al 30/6).
    """
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, data_only=True)["Movimientos"]
    out = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or str(r[2]).strip().lower() != "banco":
            continue
        f = r[0].date() if isinstance(r[0], datetime.datetime) else None
        if f:
            out.add((f.year, f.month, round(float(r[5] or 0), 2)))
    return out


# Cargos chicos del banco que NO se cargan uno por uno: se agrupan en una sola
# línea al último día del mes. Criterio verificado contra junio 2026, que cerró
# al centavo. patrón -> (etiqueta, categoría)
AGRUPADOS = [
    (r"SIRCREB",        "IIBB SIRCREB",              "Ingresos Brutos"),
    (r"DBCR 25413",     "Impuesto Ley 25413",        "Gastos bancarios"),
    (r"COMISION",       "Comisiones transferencias", "Gastos bancarios"),
]


def agrupar(movs):
    """Colapsa los cargos bancarios chicos en una línea mensual cada uno."""
    sueltos, bolsas = [], {}
    for m in movs:
        d = sin_tildes(m["obs"]).upper()
        for patron, etiqueta, categoria in AGRUPADOS:
            if re.search(patron, d):
                clave = (m["fecha"].year, m["fecha"].month, etiqueta)
                if clave not in bolsas:
                    bolsas[clave] = {
                        "fecha": m["fecha"], "tipo": "Egreso", "local": None,
                        "categoria": categoria, "monto": 0.0,
                        "obs": f"{etiqueta} {MESES[m['fecha'].month]} "
                               f"{m['fecha'].year} (extracto)",
                    }
                b = bolsas[clave]
                b["monto"] += m["monto"]
                b["fecha"] = max(b["fecha"], m["fecha"])   # queda al cierre
                break
        else:
            sueltos.append(m)
    return sueltos + sorted(bolsas.values(), key=lambda x: x["obs"])


# ---------------------------------------------------------------------------
# Cruce contra las facturas de la carpeta del mes
# ---------------------------------------------------------------------------
# Cada proveedor escribe el importe distinto: AFIP saca «1936000,00» sin
# separador de miles, Andersen imprime «$ 599,251.55» en formato yanqui y las
# liquidaciones usan «1.234,56». Se captura crudo y se normaliza abajo.
NUM = r"[\d][\d.,]{2,18}[\d]"
ETIQ = r"(?:TOTAL|Importe Total)"
# La etiqueta puede ir ANTES del número (AFIP) o DESPUÉS (Andersen: el texto
# sale «… NETO: $ 599,251.55 TOTAL:» porque la tabla se extrae por columnas).
# Se guardan TODOS los candidatos y gana el que matchee un débito del extracto
# al peso: un candidato de más no inventa nada, sólo no matchea.
TOTAL_PDF = re.compile(rf"{ETIQ}[^0-9\-]{{0,40}}({NUM})", re.I)
TOTAL_PDF_ATRAS = re.compile(rf"({NUM})[^0-9\-]{{0,20}}{ETIQ}", re.I)
CUIT_PDF = re.compile(r"\b(\d{11})\b")
# Comprobantes de pago del Macro (subcarpeta «Pagos/»). Traen el mismo «Nro. de
# Referencia» de 8 dígitos que el banco después imprime en la glosa del débito
# («N/D Transf. MacrOnline E-set D/T 87135856»), así que el cruce es por
# IDENTIDAD y no por importe. Hace falta porque esos débitos salen sin CUIT y
# porque el importe del comprobante NO es el del extracto: el «Importe total»
# suma $100 de comisión + $21 de IVA que el banco debita en un renglón aparte.
# Se usa el «IMPORTE A TRANSFERIR», que sí coincide al centavo.
REF_PAGO = re.compile(r"(\d{8})\s*\n\s*Nro\. de Referencia")
IMPORTE_TRANSFERIDO = re.compile(r"IMPORTE A TRANSFERIR\s*\n\s*\$\s*([\d.,]+)")
# El CUIT del EMISOR, no el nuestro: Mahni aparece en todas las facturas.
CUIT_PROPIO = "30719012503"


def plata_libre(s):
    """Un importe escrito en cualquiera de los tres formatos -> float.

    El separador DECIMAL es el último punto o coma que quede seguido de
    exactamente dos dígitos al final. Todo lo anterior es separador de miles y
    se tira. Sin esto, «599,251.55» se lee como $599,25.
    """
    s = s.strip()
    if not re.search(r"[.,]\d{2}$", s):
        return float(re.sub(r"[.,]", "", s))
    entero, dec = s[:-3], s[-2:]
    return float(re.sub(r"[.,]", "", entero) + "." + dec)


def indexar_facturas(carpeta):
    """[(cuit, nombre_archivo, total)] de cada PDF de la carpeta.

    El CUIT sale primero del nombre del archivo (convención de AFIP) y, si el
    archivo no la respeta, del texto. Un PDF sin total legible se saltea con
    aviso: mejor que un match inventado.
    """
    facturas, ilegibles = [], []
    # Recursivo: los comprobantes de pago viven en la subcarpeta «Pagos/» y son
    # los únicos que identifican los débitos «N/D Transf. MacrOnline», que salen
    # del banco sin CUIT en la glosa. El importe pagado casi nunca coincide con
    # una factura del mes (se pagan facturas viejas), pero sí con su comprobante.
    for ruta in sorted(glob.glob(os.path.join(carpeta, "**", "*.pdf"), recursive=True)):
        base = os.path.relpath(ruta, carpeta)
        try:
            texto = fitz.open(ruta)[0].get_text()
        except Exception as e:
            ilegibles.append(f"{base}: no se pudo abrir ({e})")
            continue
        m = re.match(r"(\d{11})_", base)
        cuit = m.group(1) if m else next(
            (c for c in CUIT_PDF.findall(texto) if c != CUIT_PROPIO), None)
        # Comprobante de pago del Macro: manda la referencia, y el importe bueno
        # es el transferido (sin la comisión que el banco debita por separado).
        ref = REF_PAGO.search(texto)
        transferido = IMPORTE_TRANSFERIDO.search(texto)
        if ref and transferido:
            facturas.append((cuit, base, [plata_libre(transferido.group(1))],
                             ref.group(1)))
            continue
        totales = sorted({plata_libre(t) for t in
                          TOTAL_PDF.findall(texto) + TOTAL_PDF_ATRAS.findall(texto)})
        if not totales:
            ilegibles.append(f"{base}: no le encontré el total")
            continue
        facturas.append((cuit, base, totales, None))
    return facturas, ilegibles


def resolver_con_facturas(movs, facturas):
    """Le pone categoría a los egresos que la glosa del banco no identifica.

    Primero por «Nro. de Referencia»: los comprobantes de pago del Macro traen
    el mismo número de 8 dígitos que el banco imprime en la glosa del débito,
    así que el cruce es exacto y no depende del importe. El importe se sigue
    exigiendo como control: si la referencia matchea pero la plata no, no se
    resuelve solo — eso es algo para mirar.

    Si no hay referencia, cae al match por importe con tolerancia de $1: el
    banco redondea al peso (Andersen facturó $599.251,55 y el débito salió
    $599.251,00).

    Una factura paga UN débito: en septiembre 2026 Annie ($1.500.000 el 15/09) y
    una transferencia a otra cuenta de la propia MAHNI (mismo importe, el 16/09)
    se llevaban la misma factura y el segundo débito quedaba como Marketing. Y si
    la glosa trae un CUIT, la factura tiene que ser de ese CUIT.
    """
    resueltos, usadas = [], set()
    for m in movs:
        if m["local"] or m["categoria"] or m["tipo"] != "Egreso":
            continue
        cuits_glosa = set(re.findall(r"(?<!\d)(\d{11})(?!\d)", m["obs"]))
        disponibles = [f for f in facturas if f[1] not in usadas
                       and (not cuits_glosa or f[0] in cuits_glosa)]
        cand = [f for f in disponibles
                if f[3] and re.search(rf"\b{f[3]}\b", m["obs"])]
        if len(cand) == 1 and not any(abs(t - m["monto"]) < 1.0 for t in cand[0][2]):
            m["ambiguo"] = [f"{cand[0][1]} (referencia {cand[0][3]} coincide pero "
                            f"el importe no: comprobante {cand[0][2][0]:,.2f} vs "
                            f"extracto {m['monto']:,.2f})"]
            continue
        if not cand:
            cand = [f for f in disponibles
                    if any(abs(t - m["monto"]) < 1.0 for t in f[2])]
        if len(cand) != 1:
            if len(cand) > 1:
                m["ambiguo"] = [c[1] for c in cand]
            continue
        cuit, archivo, totales, _ref = cand[0]
        total = min(totales, key=lambda t: abs(t - m["monto"]))
        if cuit not in CUIT_CATEGORIA:
            m["ambiguo"] = [f"{archivo} (CUIT {cuit} no está en CUIT_CATEGORIA)"]
            continue
        nombre, categoria = CUIT_CATEGORIA[cuit]
        m["categoria"] = categoria
        usadas.add(archivo)
        m["obs"] = f"{m['obs']} — {nombre} {archivo.replace('.pdf', '')}"
        resueltos.append((m, nombre, total))
    return resueltos


# ---------------------------------------------------------------------------
# La planilla
# ---------------------------------------------------------------------------
def movimientos_del_sheet():
    from google_auth import sheets
    return sheets().spreadsheets().values().get(
        spreadsheetId=MASTER_PLAN, range="Movimientos!A1:H2000",
        valueRenderOption="UNFORMATTED_VALUE").execute().get("values", [])


def ya_cargados_sheet(filas):
    """(año, mes, monto) de lo que ya está cargado con Medio «Banco»."""
    out = set()
    for r in filas[1:]:
        r = list(r) + [""] * (8 - len(r))
        if str(r[2]).strip().lower() != "banco":
            continue
        try:
            f = datetime.date(1899, 12, 30) + datetime.timedelta(days=float(r[0]))
            out.add((f.year, f.month, round(float(r[5] or 0), 2)))
        except (ValueError, TypeError):
            continue
    return out


def escribir(filas_nuevas, cantidad_actual):
    """Appendea al final de Movimientos y RELEE para confirmar."""
    from google_auth import sheets
    s = sheets()
    s.spreadsheets().values().append(
        spreadsheetId=MASTER_PLAN, range="Movimientos!A:H",
        valueInputOption="USER_ENTERED", insertDataOption="INSERT_ROWS",
        body={"values": filas_nuevas}).execute()
    despues = movimientos_del_sheet()
    agregadas = len(despues) - cantidad_actual
    if agregadas != len(filas_nuevas):
        sys.exit(f"ERROR: quise escribir {len(filas_nuevas)} filas y la planilla "
                 f"creció {agregadas}. Revisá Movimientos a mano.")
    return agregadas


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pdf", help="el PDF del extracto del Macro")
    ap.add_argument("--facturas", help="carpeta de Facturas de Compra del mes")
    ap.add_argument("--escribir", action="store_true",
                    help="escribe en Movimientos (por defecto NO toca nada)")
    ap.add_argument("--forzar", action="store_true",
                    help="escribe aunque queden renglones sin categoría")
    ap.add_argument("--saltear-sin-categoria", action="store_true",
                    help="escribe SÓLO los renglones resueltos y deja afuera los "
                         "que no tienen categoría (los lista al final). Para "
                         "cuando falta identificar un pago y no se quiere frenar "
                         "el resto del mes: los salteados entran después, y el "
                         "dedupe por (año, mes, monto) evita cargarlos dos veces")
    args = ap.parse_args()

    crudos = parsear(args.pdf)
    movs = agrupar(crudos)

    # El neto se controla contra el extracto CRUDO, no contra el agrupado.
    ing = sum(m["monto"] for m in crudos if m["tipo"] == "Ingreso")
    egr = sum(m["monto"] for m in crudos if m["tipo"] == "Egreso")
    print(f"# {len(movs)} movimientos en el extracto — "
          f"ingresos ${ing:,.2f} · egresos ${egr:,.2f} · neto ${ing - egr:,.2f}")

    if args.facturas:
        facturas, ilegibles = indexar_facturas(args.facturas)
        resueltos = resolver_con_facturas(movs, facturas)
        print(f"# {len(facturas)} facturas indexadas en la carpeta, "
              f"{len(resueltos)} débitos identificados por importe")
        for m, nombre, total in resueltos:
            print(f"#   ${m['monto']:,.2f} → {nombre} (la factura dice "
                  f"${total:,.2f})")
        for x in ilegibles:
            print(f"#   ⚠ {x}")

    filas_sheet = movimientos_del_sheet()
    cargados = ya_cargados_sheet(filas_sheet)
    nuevos = [m for m in movs
              if (m["fecha"].year, m["fecha"].month, round(m["monto"], 2))
              not in cargados]
    print(f"# {len(movs) - len(nuevos)} ya estaban cargados, "
          f"{len(nuevos)} para agregar")
    print("#")

    filas = []
    sin_resolver = []
    for m in nuevos:
        falta = not (m["local"] or m["categoria"])
        if falta:
            sin_resolver.append(m)
        filas.append([
            m["fecha"].strftime("%-d/%-m/%Y"),
            m["tipo"], "Banco",
            m["local"] or "", m["categoria"] or "",
            round(m["monto"], 2), "ARS",
            f"{m['obs']} - {MESES[m['fecha'].month]} {m['fecha'].year}",
        ])
        marca = "  ← SIN CATEGORIA" if falta else ""
        if m.get("ambiguo"):
            marca = f"  ← {len(m['ambiguo'])} facturas posibles: {m['ambiguo']}"
        print(f"{filas[-1][0]}\t{m['tipo']}\tBanco\t{m['local'] or ''}\t"
              f"{m['categoria'] or ''}\t{m['monto']:.2f}\tARS\t"
              f"{filas[-1][7]}{marca}")

    if sin_resolver and args.saltear_sin_categoria:
        sin = {id(m) for m in sin_resolver}
        filas = [f for m, f in zip(nuevos, filas) if id(m) not in sin]
        print(f"\n[--saltear-sin-categoria] Dejo AFUERA {len(sin_resolver)} "
              f"movimiento(s) sin identificar, por ${sum(m['monto'] for m in sin_resolver):,.2f}:")
        for m in sin_resolver:
            print(f"    {m['fecha'].strftime('%-d/%-m/%Y')}  ${m['monto']:,.2f}  {m['obs']}")
        print("  Cuando se sepan qué son, se vuelve a correr y entran solos.")
    if not args.escribir:
        print(f"\n[SIN --escribir] No toqué nada. {len(filas)} filas listas.")
        return
    if sin_resolver and not args.forzar and not args.saltear_sin_categoria:
        sys.exit(f"\nFRENO: {len(sin_resolver)} movimiento(s) sin categoría. "
                 f"Entrarían a Movimientos y desaparecerían de todos los SUMIFS "
                 f"sin fallar. Resolvelos (agregá el CUIT a CUIT_CATEGORIA o la "
                 f"glosa a REGLAS) o corré con --forzar si de verdad van vacíos.")
    if not filas:
        print("\nNada nuevo para escribir.")
        return
    n = escribir(filas, len(filas_sheet))
    print(f"\n✅ {n} filas escritas en Movimientos y verificadas releyendo.")


if __name__ == "__main__":
    main()
